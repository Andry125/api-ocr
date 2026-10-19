import io
import os

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from google import genai
from google.genai import types
from google.genai.errors import APIError
from PIL import Image
from pydantic import BaseModel, Field
from tenacity import retry, retry_if_exception, stop_after_attempt, wait_exponential

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.0-flash")
client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

app = FastAPI(title="API OCR Gemini — texte, JSON structuré, Word")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Schémas Pydantic — servent de response_schema à Gemini ET de validation
# ---------------------------------------------------------------------------
class TableMerge(BaseModel):
    row: int = Field(description="Index 0-based de la ligne de départ")
    col: int = Field(description="Index 0-based de la colonne de départ")
    rowspan: int = 1
    colspan: int = 1


class DocTable(BaseModel):
    title: str = ""
    headers: list[str] = Field(default_factory=list)
    rows: list[list[str]] = Field(default_factory=list)
    merges: list[TableMerge] = Field(default_factory=list)


class DocBlock(BaseModel):
    type: str = Field(
        description="Type de bloc : 'heading', 'paragraph', 'list_item' ou 'table'"
    )
    level: int = Field(0, description="Niveau de titre (1-6) si type='heading'")
    text: str = Field("", description="Contenu texte du bloc (hors tableaux)")
    table: DocTable = Field(default_factory=DocTable)


class DocumentExtraction(BaseModel):
    blocks: list[DocBlock] = Field(default_factory=list)


# ---------------------------------------------------------------------------
# Retry (identique à votre version)
# ---------------------------------------------------------------------------
def is_transient_error(exception):
    if isinstance(exception, APIError):
        code = getattr(exception, "code", None)
        err_str = str(exception)
        return code in [503, 429] or "503" in err_str or "429" in err_str
    return False


@retry(
    stop=stop_after_attempt(4),
    wait=wait_exponential(multiplier=1, min=1, max=8),
    retry=retry_if_exception(is_transient_error),
    reraise=True,
)
def call_gemini_with_retry(model_name: str, contents: list, config=None):
    kwargs = {"model": model_name, "contents": contents}
    if config is not None:
        kwargs["config"] = config
    return client.models.generate_content(**kwargs)


# ---------------------------------------------------------------------------
# Prompts
# ---------------------------------------------------------------------------
PROMPT_TEXT = (
    "Extrais tout le texte visible dans cette image. "
    "Restitue uniquement le texte extrait, sans aucun commentaire ni explication. "
    "Conserve et reproduis les structures de tableaux : en-têtes, lignes, colonnes, "
    "cellules vides et organisation des données. "
    "Représente les tableaux au format Markdown ; si un tableau ne peut pas être "
    "représenté correctement en Markdown, utilise HTML. "
    "Préserve l'ordre de lecture et les sauts de ligne utiles."
)

PROMPT_STRUCTURED = (
    "Analyse cette image et reconstruis sa structure logique de document. "
    "Renvoie une liste ordonnée de blocs dans 'blocks'. "
    "Pour chaque titre, utilise type='heading' avec 'level' entre 1 et 6. "
    "Pour chaque paragraphe, type='paragraph'. "
    "Pour chaque élément de liste, type='list_item'. "
    "Pour chaque tableau, type='table' et remplis 'table.headers', 'table.rows' "
    "et, si des cellules sont fusionnées, 'table.merges' "
    "(row/col 0-based, rowspan/colspan). "
    "Conserve les cellules vides sous forme de chaîne vide. "
    "N'ajoute aucun commentaire en dehors du JSON."
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _ensure_client():
    if not client:
        raise HTTPException(500, "La clé GEMINI_API_KEY n'est pas configurée.")


def _validate_image(file: UploadFile):
    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(400, "Le fichier fourni doit être une image.")


async def _extract_text(file: UploadFile) -> dict:
    _ensure_client()
    _validate_image(file)
    try:
        image = Image.open(file.file)
        response = call_gemini_with_retry(
            model_name=GEMINI_MODEL, contents=[PROMPT_TEXT, image]
        )
        return {
            "filename": file.filename,
            "extracted_text": response.text.strip() if response.text else "",
        }
    except APIError as e:
        raise HTTPException(
            502, f"Gemini indisponible après plusieurs tentatives : {e}"
        )
    except Exception as e:
        raise HTTPException(500, f"Erreur lors du traitement : {e}")


async def _extract_structured(file: UploadFile) -> DocumentExtraction:
    _ensure_client()
    _validate_image(file)
    try:
        image = Image.open(file.file)
        config = types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=DocumentExtraction,
            temperature=0,
        )
        response = call_gemini_with_retry(
            model_name=GEMINI_MODEL,
            contents=[PROMPT_STRUCTURED, image],
            config=config,
        )
        raw = response.text or "{}"
        return DocumentExtraction.model_validate_json(raw)
    except APIError as e:
        raise HTTPException(
            502, f"Gemini indisponible après plusieurs tentatives : {e}"
        )
    except Exception as e:
        raise HTTPException(500, f"Erreur lors du traitement : {e}")


# ---------------------------------------------------------------------------
# Reconstruction Word
# ---------------------------------------------------------------------------
def _append_table(doc: Document, t: DocTable):
    n_cols = max([len(t.headers)] + [len(r) for r in t.rows] + [1])
    n_body_rows = len(t.rows)
    n_total_rows = n_body_rows + (1 if t.headers else 0)
    table = doc.add_table(rows=n_total_rows, cols=n_cols)
    table.style = "Table Grid"

    offset = 0
    if t.headers:
        for j, h in enumerate(t.headers):
            if j < n_cols:
                table.rows[0].cells[j].text = h
        offset = 1

    for i, row in enumerate(t.rows):
        for j, cell_text in enumerate(row):
            if j < n_cols:
                table.rows[i + offset].cells[j].text = cell_text

    for m in t.merges:
        r = m.row + offset
        c = m.col
        rs = max(1, m.rowspan)
        cs = max(1, m.colspan)
        if rs == 1 and cs == 1:
            continue
        r2 = min(r + rs - 1, n_total_rows - 1)
        c2 = min(c + cs - 1, n_cols - 1)
        if r2 > r or c2 > c:
            table.cell(r, c).merge(table.cell(r2, c2))


def build_docx(extraction: DocumentExtraction) -> bytes:
    doc = Document()
    for block in extraction.blocks:
        if block.type == "heading":
            level = block.level if 1 <= block.level <= 9 else 1
            doc.add_heading(block.text, level=level)
        elif block.type == "list_item":
            doc.add_paragraph(block.text, style="List Bullet")
        elif block.type == "table":
            _append_table(doc, block.table)
        else:  # paragraph par défaut
            doc.add_paragraph(block.text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------
@app.post("/ocr")
async def extract_text_from_image(file: UploadFile = File(...)):
    return await _extract_text(file)


@app.post("/ocr/structured")
async def extract_structured(file: UploadFile = File(...)):
    result = await _extract_structured(file)
    return result.model_dump()


@app.post("/ocr/docx")
async def extract_to_docx(file: UploadFile = File(...)):
    result = await _extract_structured(file)
    data = build_docx(result)
    stem = (file.filename or "document").rsplit(".", 1)[0]
    return StreamingResponse(
        io.BytesIO(data),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={"Content-Disposition": f'attachment; filename="{stem}.docx"'},
    )
